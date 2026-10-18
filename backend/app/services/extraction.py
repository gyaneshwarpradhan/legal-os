"""Text extraction pipeline with multi-format parsing and OCR abstraction."""
import io
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional


@dataclass
class ExtractedPage:
    page_number: int
    text: str
    needs_ocr: bool = False
    ocr_confidence: Optional[float] = None


@dataclass
class ExtractedDocument:
    file_name: str
    file_type: str
    total_pages: int
    pages: List[ExtractedPage]
    full_text: str


class OCRProvider(ABC):
    @abstractmethod
    def extract_text_from_image(self, image_bytes: bytes) -> str:
        """Run optical character recognition on image/scanned bytes."""
        pass


class MockLocalOCRProvider(OCRProvider):
    """Fallback OCR provider for development and testing."""

    def extract_text_from_image(self, image_bytes: bytes) -> str:
        return "[OCR Extracted Text from scanned legal document]"


class AWSTextractOCRProvider(OCRProvider):
    """Production AWS Textract OCR provider adapter."""

    def __init__(self, region_name: str = "us-east-1"):
        self.region_name = region_name

    def extract_text_from_image(self, image_bytes: bytes) -> str:
        try:
            import boto3

            client = boto3.client("textract", region_name=self.region_name)
            response = client.detect_document_text(
                Document={"Bytes": image_bytes}
            )
            lines = [
                item["Text"]
                for item in response.get("Blocks", [])
                if item["BlockType"] == "LINE"
            ]
            return "\n".join(lines)
        except Exception as e:
            return f"[Textract Error: {str(e)}]"


class DocumentExtractor:
    def __init__(self, ocr_provider: Optional[OCRProvider] = None):
        self.ocr_provider = ocr_provider or MockLocalOCRProvider()

    def extract(self, file_name: str, file_bytes: bytes) -> ExtractedDocument:
        ext = Path(file_name).suffix.lower()

        if ext == ".txt":
            return self._extract_txt(file_name, file_bytes)
        elif ext == ".csv":
            return self._extract_csv(file_name, file_bytes)
        elif ext == ".docx":
            return self._extract_docx(file_name, file_bytes)
        elif ext == ".pdf":
            return self._extract_pdf(file_name, file_bytes)
        elif ext in {".png", ".jpg", ".jpeg"}:
            return self._extract_image(file_name, file_bytes)
        else:
            # Fallback to UTF-8 text decoding
            text = file_bytes.decode("utf-8", errors="replace")
            page = ExtractedPage(page_number=1, text=text)
            return ExtractedDocument(
                file_name=file_name,
                file_type=ext,
                total_pages=1,
                pages=[page],
                full_text=text,
            )

    def _extract_txt(self, file_name: str, file_bytes: bytes) -> ExtractedDocument:
        text = file_bytes.decode("utf-8", errors="replace")
        page = ExtractedPage(page_number=1, text=text)
        return ExtractedDocument(
            file_name=file_name,
            file_type="txt",
            total_pages=1,
            pages=[page],
            full_text=text,
        )

    def _extract_csv(self, file_name: str, file_bytes: bytes) -> ExtractedDocument:
        text = file_bytes.decode("utf-8", errors="replace")
        page = ExtractedPage(page_number=1, text=text)
        return ExtractedDocument(
            file_name=file_name,
            file_type="csv",
            total_pages=1,
            pages=[page],
            full_text=text,
        )

    def _extract_docx(self, file_name: str, file_bytes: bytes) -> ExtractedDocument:
        try:
            import docx

            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n\n".join(paragraphs)
            page = ExtractedPage(page_number=1, text=full_text)
            return ExtractedDocument(
                file_name=file_name,
                file_type="docx",
                total_pages=1,
                pages=[page],
                full_text=full_text,
            )
        except Exception:
            text = file_bytes.decode("utf-8", errors="replace")
            page = ExtractedPage(page_number=1, text=text)
            return ExtractedDocument(
                file_name=file_name,
                file_type="docx",
                total_pages=1,
                pages=[page],
                full_text=text,
            )

    def _extract_pdf(self, file_name: str, file_bytes: bytes) -> ExtractedDocument:
        try:
            import pypdf

            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            pages = []
            full_text_parts = []

            for idx, p in enumerate(reader.pages):
                page_num = idx + 1
                page_text = p.extract_text() or ""
                needs_ocr = len(page_text.strip()) < 30

                if needs_ocr:
                    # In a scanned page, request OCR from provider
                    ocr_text = self.ocr_provider.extract_text_from_image(file_bytes)
                    page_text = page_text + ("\n" if page_text else "") + ocr_text

                pages.append(
                    ExtractedPage(
                        page_number=page_num,
                        text=page_text,
                        needs_ocr=needs_ocr,
                    )
                )
                full_text_parts.append(page_text)

            return ExtractedDocument(
                file_name=file_name,
                file_type="pdf",
                total_pages=len(pages),
                pages=pages,
                full_text="\n\n".join(full_text_parts),
            )
        except Exception:
            text = file_bytes.decode("utf-8", errors="replace")
            page = ExtractedPage(page_number=1, text=text)
            return ExtractedDocument(
                file_name=file_name,
                file_type="pdf",
                total_pages=1,
                pages=[page],
                full_text=text,
            )

    def _extract_image(self, file_name: str, file_bytes: bytes) -> ExtractedDocument:
        ocr_text = self.ocr_provider.extract_text_from_image(file_bytes)
        page = ExtractedPage(page_number=1, text=ocr_text, needs_ocr=True)
        return ExtractedDocument(
            file_name=file_name,
            file_type="image",
            total_pages=1,
            pages=[page],
            full_text=ocr_text,
        )


extractor = DocumentExtractor()
